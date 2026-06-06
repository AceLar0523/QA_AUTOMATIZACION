from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

def set_cell_background(cell, color_hex):
    """Set background color for a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def create_test_report(test_data, filepath):
    """
    Generates a Word Document based on the provided test data and template.
    test_data dictionary expects:
    - title
    - test_id
    - date (YYYY-MM-DD)
    - author (e.g. 'Agente IA')
    - use_case
    - module
    - version
    - description
    - required_data
    - prerequisites
    - postconditions
    - notes
    - steps: list of dicts:
        {'step_no': 1, 'action': '...', 'expected': '...', 'obtained': '...', 'defects': '', 'status': 'PASA'}
    """
    doc = Document()
    
    # Header
    p = doc.add_paragraph()
    r = p.add_run(f"Test: {test_data['title']}")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 51, 102)
    
    now = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    p2 = doc.add_paragraph(f"Fecha de ejecución técnica: {now}")
    p2.style.font.size = Pt(10)
    
    # First Table (Metadata)
    table1 = doc.add_table(rows=7, cols=2)
    table1.style = 'Table Grid'
    
    table1.cell(0, 0).text = f"Título: {test_data['title']}"
    table1.cell(0, 1).text = f"ID: {test_data['test_id']}"
    table1.cell(1, 0).text = f"Autor: {test_data['author']}"
    table1.cell(1, 1).text = f"Fecha: {test_data['date']}"
    
    cell_uc = table1.cell(2, 0)
    cell_uc.merge(table1.cell(2, 1))
    cell_uc.text = f"Caso de Uso perteneciente: {test_data['use_case']}"
    
    table1.cell(3, 0).text = f"Módulo: {test_data['module']}"
    table1.cell(3, 1).text = f"Versión del Sistema: {test_data['version']}"
    
    cell_desc = table1.cell(4, 0)
    cell_desc.merge(table1.cell(4, 1))
    cell_desc.text = f"Descripción: {test_data['description']}"
    
    table1.cell(5, 0).text = f"Datos Requeridos: {test_data['required_data']}"
    table1.cell(5, 1).text = f"Prerrequisitos: {test_data['prerequisites']}"
    
    cell_post = table1.cell(6, 0)
    cell_post.merge(table1.cell(6, 1))
    cell_post.text = f"Postcondiciones: {test_data['postconditions']}"
    
    doc.add_paragraph() # Spacer
    doc.add_paragraph(f"Notas: {test_data['notes']}")
    doc.add_paragraph() # Spacer
    
    # Second Table (Execution Steps)
    table2 = doc.add_table(rows=4 + len(test_data['steps']), cols=6)
    table2.style = 'Table Grid'
    
    table2.cell(0, 0).merge(table2.cell(0, 5)).text = f"ID: {test_data['test_id']} | Fecha: {test_data['date']}"
    table2.cell(1, 0).merge(table2.cell(1, 5)).text = f"Responsable: {test_data['author']}"
    
    headers = ["PASO", "ENTRADA O ACCIÓN", "RESULTADO ESPERADO", "RESULTADO OBTENIDO", "DEFECTOS", "PASA / FALLA"]
    for i, h in enumerate(headers):
        table2.cell(2, i).text = h
        
    table2.cell(3, 0).merge(table2.cell(3, 5)).text = "FLUJO A PROBAR"
    
    row_idx = 4
    for step in test_data['steps']:
        table2.cell(row_idx, 0).text = str(step['step_no'])
        table2.cell(row_idx, 1).text = step['action']
        table2.cell(row_idx, 2).text = step['expected']
        table2.cell(row_idx, 3).text = step['obtained']
        table2.cell(row_idx, 4).text = step['defects']
        table2.cell(row_idx, 5).text = step['status']
        row_idx += 1
        
    doc.add_paragraph() # Spacer
    
    # Third Table (Evidence)
    table3 = doc.add_table(rows=4 + len(test_data['steps']), cols=4)
    table3.style = 'Table Grid'
    
    table3.cell(0, 0).merge(table3.cell(0, 3)).text = f"ID: {test_data['test_id']} | Fecha: {test_data['date']}"
    table3.cell(1, 0).merge(table3.cell(1, 3)).text = f"Responsable: {test_data['author']}"
    
    headers_evi = ["PASO", "ENTRADA O ACCIÓN", "OBSERVACIONES", "EVIDENCIA"]
    for i, h in enumerate(headers_evi):
        table3.cell(2, i).text = h
        
    table3.cell(3, 0).merge(table3.cell(3, 3)).text = "FLUJO A PROBAR"
    from docx.shared import Inches
    
    row_idx = 4
    for step in test_data['steps']:
        table3.cell(row_idx, 0).text = str(step['step_no'])
        table3.cell(row_idx, 1).text = step['action']
        
        # Llenar observación
        table3.cell(row_idx, 2).text = step.get('observations', 'Validación completada')
        
        # Llenar evidencia (imagen)
        cell_evidence = table3.cell(row_idx, 3)
        screenshot_path = step.get('screenshot_path')
        if screenshot_path:
            p = cell_evidence.paragraphs[0]
            run = p.add_run()
            try:
                run.add_picture(screenshot_path, width=Inches(2.0))
            except Exception as e:
                cell_evidence.text = f"Error al cargar imagen: {e}"
                
        row_idx += 1
        
    doc.save(filepath)
    print(f"Reporte generado exitosamente en: {filepath}")
